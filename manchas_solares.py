"""Experimento AR(p) anual com treino, validação e teste cronológicos."""
import hashlib
import json
from pathlib import Path

import numpy as np
import pandas as pd
from sklearn.metrics import mean_absolute_error, mean_squared_error, r2_score
from statsmodels.tsa.ar_model import AutoReg
from openpyxl.styles import Alignment, Font, PatternFill

MAX_LAGS = 30


def carregar(pasta):
    caminho = Path(pasta) / "dados_manchas_solares" / "sunspots.csv"
    fonte = json.loads(caminho.with_name("fonte.json").read_text())
    if hashlib.sha256(caminho.read_bytes()).hexdigest() != fonte["sha256"]:
        raise ValueError("O CSV difere da cópia registrada em fonte.json.")
    dados = pd.read_csv(caminho).rename(columns={"YEAR": "Ano", "SUNACTIVITY": "Atividade solar"})
    dados["Ano"] = dados["Ano"].astype(int)
    dados["Atividade solar"] = pd.to_numeric(dados["Atividade solar"], errors="raise")
    if len(dados) != 309 or not np.array_equal(dados["Ano"], np.arange(1700, 2009)):
        raise ValueError("Esperados 309 anos consecutivos, de 1700 a 2008.")
    valores = dados["Atividade solar"].to_numpy(dtype=float)
    if not np.isfinite(valores).all() or (valores < 0).any():
        raise ValueError("Há valores inválidos na atividade solar.")
    return dados, fonte


def erro_relativo(real, previsto):
    real, previsto = np.asarray(real, dtype=float), np.asarray(previsto, dtype=float)
    resultado = np.full(real.shape, np.nan)
    np.divide(np.abs(previsto - real), np.abs(real), out=resultado,
              where=(real != 0) & np.isfinite(real) & np.isfinite(previsto))
    return resultado


def prever_um_passo(ajuste, historico):
    lags = np.asarray(ajuste.model.ar_lags, dtype=int)
    historico = np.asarray(historico, dtype=float)
    if len(historico) < lags.max():
        raise ValueError("Histórico insuficiente para as defasagens do modelo.")
    # O AutoReg foi ajustado com trend='c': intercepto seguido dos coeficientes.
    valor = float(ajuste.params[0] + np.dot(ajuste.params[1:], historico[-lags]))
    # A atividade solar não pode ser negativa. A mesma regra vale na validação e no teste.
    return max(0.0, valor)


def avaliar_sequencial(ajuste, dados, inicio, fim):
    historico = dados["Atividade solar"].iloc[:inicio].to_list()
    linhas = []
    for i in range(inicio, fim):
        previsto = prever_um_passo(ajuste, historico)
        persistencia = historico[-1]
        # O real do alvo só é acessado depois que a previsão foi produzida.
        real = float(dados["Atividade solar"].iloc[i])
        ano = int(dados["Ano"].iloc[i])
        linhas.append([ano - 1, ano, real, previsto, abs(real - previsto),
                       float(erro_relativo([real], [previsto])[0]), persistencia])
        historico.append(real)
    return pd.DataFrame(linhas, columns=[
        "Ano de origem", "Ano previsto", "Valor real", "Valor predito",
        "Erro absoluto", "Taxa de erro (%)", "Persistência"
    ])


def calcular_metricas(resultados, valores_ajuste):
    real = resultados["Valor real"].to_numpy()
    escala_mase = np.abs(np.diff(valores_ajuste)).mean()
    linhas = []
    for nome, coluna in [("Autorregressivo", "Valor predito"), ("Persistência", "Persistência")]:
        previsto = resultados[coluna].to_numpy()
        erros = erro_relativo(real, previsto)
        mae = mean_absolute_error(real, previsto)
        linhas.append({
            "Modelo": nome,
            "Anos avaliados": len(real),
            "MAE": mae,
            "RMSE": np.sqrt(mean_squared_error(real, previsto)),
            "R²": r2_score(real, previsto),
            "MAPE (%)": np.nanmean(erros) if np.isfinite(erros).any() else np.nan,
            "MASE": mae / escala_mase if escala_mase > 0 else np.nan,
        })
    return pd.DataFrame(linhas)


def executar(pasta):
    dados, fonte = carregar(pasta)
    n = len(dados)
    fim_treino = int(n * 0.60)
    fim_validacao = int(n * 0.80)
    # Cortes de 60% e 80%: 185 de treino, 62 de validação e 62 de teste.
    valores = dados["Atividade solar"].to_numpy()
    selecao = []
    for p in range(1, MAX_LAGS + 1):
        # Mesmo primeiro alvo de treino para comparar todas as ordens de forma consistente.
        ajuste = AutoReg(valores[:fim_treino], lags=p, trend="c", hold_back=MAX_LAGS,
                        ).fit()
        validacao = avaliar_sequencial(ajuste, dados, fim_treino, fim_validacao)
        selecao.append({
            "Defasagens (p)": p,
            "MAE validação": mean_absolute_error(validacao["Valor real"], validacao["Valor predito"]),
            "RMSE validação": np.sqrt(mean_squared_error(validacao["Valor real"], validacao["Valor predito"])),
        })
    selecao = pd.DataFrame(selecao).sort_values(["MAE validação", "Defasagens (p)"]).reset_index(drop=True)
    p = int(selecao.iloc[0]["Defasagens (p)"])
    # A ordem é congelada antes de consultar qualquer resultado do teste.
    ajuste_validacao = AutoReg(valores[:fim_treino], lags=p, trend="c", hold_back=MAX_LAGS,
                              ).fit()
    validacao = avaliar_sequencial(ajuste_validacao, dados, fim_treino, fim_validacao)
    # Após selecionar p, ajustamos os coeficientes com treino + validação.
    ajuste_final = AutoReg(valores[:fim_validacao], lags=p, trend="c").fit()
    teste = avaliar_sequencial(ajuste_final, dados, fim_validacao, n)
    metricas = calcular_metricas(teste, valores[:fim_validacao])
    dados = dados.copy()
    dados["Partição"] = np.where(np.arange(n) < fim_treino, "Treino",
                                np.where(np.arange(n) < fim_validacao, "Validação", "Teste"))
    coeficientes = pd.DataFrame({
        "Parâmetro": ["Intercepto"] + [f"Lag {i}" for i in range(1, p + 1)],
        "Coeficiente": ajuste_final.params,
    })
    configuracao = pd.DataFrame([
        ["Dataset", "Sunspots anual, 309 observações de 1700 a 2008"],
        ["Fonte", fonte["url_csv"]],
        ["Documentação", fonte["documentacao"]],
        ["SHA256", fonte["sha256"]],
        ["Treino inicial", f"{dados['Ano'].iloc[0]}–{dados['Ano'].iloc[fim_treino-1]} ({fim_treino} anos)"],
        ["Validação", f"{dados['Ano'].iloc[fim_treino]}–{dados['Ano'].iloc[fim_validacao-1]} ({fim_validacao-fim_treino} anos)"],
        ["Teste final", f"{dados['Ano'].iloc[fim_validacao]}–{dados['Ano'].iloc[-1]} ({n-fim_validacao} anos)"],
        ["Modelo", f"AR({p}), com intercepto, sem variáveis externas"],
        ["Seleção de p", f"Menor MAE na validação entre p=1 e p={MAX_LAGS}; desempate pelo menor p"],
        ["Hold-back na seleção", f"{MAX_LAGS} observações em todos os candidatos"],
        ["Ajuste final", "Treino + validação; coeficientes fixos durante o teste"],
        ["Horizonte", "Próximo ano, um passo à frente, usando apenas os valores reais dos anos anteriores"],
        ["Valores negativos previstos", "Limitados a zero na validação e no teste"],
        ["Taxa de erro", "abs(predito - real) / abs(real), exibida como %; real zero fica vazio"],
        ["MAPE", "Média do erro percentual, excluindo real zero; pode crescer perto de zero"],
        ["MASE", "MAE dividido pela média de abs(y[t]-y[t-1]) de treino + validação"],
        ["Protocolo", "Cada novo real só fica disponível depois de produzir a previsão de seu ano"],
    ], columns=["Item", "Valor"])
    return dict(dados=dados, fonte=fonte, p=p, ajuste=ajuste_final, selecao=selecao,
                validacao=validacao, teste=teste, metricas=metricas,
                coeficientes=coeficientes, configuracao=configuracao,
                fim_treino=fim_treino, fim_validacao=fim_validacao)


def exportar(e, pasta):
    destino = Path(pasta) / "resultados_manchas_solares"
    destino.mkdir(exist_ok=True)
    saida = destino / "previsao_manchas_solares.xlsx"
    tabelas = {
        "Previsoes_teste": e["teste"], "Metricas_teste": e["metricas"],
        "Selecao_defasagens": e["selecao"], "Validacao": e["validacao"],
        "Coeficientes": e["coeficientes"], "Dados": e["dados"],
        "Configuracao": e["configuracao"],
    }
    with pd.ExcelWriter(saida, engine="openpyxl") as writer:
        for nome, tabela in tabelas.items():
            tabela.to_excel(writer, sheet_name=nome, index=False)
            sheet = writer.sheets[nome]
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            sheet.sheet_view.showGridLines = False
            sheet.row_dimensions[1].height = 28
            for celula in sheet[1]:
                celula.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
                celula.fill = PatternFill("solid", fgColor="243B53")
                celula.alignment = Alignment(horizontal="center", vertical="center")
                sheet.column_dimensions[celula.column_letter].width = 25
            for linha in sheet.iter_rows(min_row=2):
                for celula in linha:
                    celula.font = Font(name="Arial", size=10)
                    celula.alignment = Alignment(vertical="center")
                    if isinstance(celula.value, float): celula.number_format = "0.000"
            if nome in ["Previsoes_teste", "Validacao"]:
                for linha in sheet.iter_rows(min_row=2): linha[5].number_format = "0.00%"
            elif nome == "Metricas_teste":
                for linha in sheet.iter_rows(min_row=2): linha[5].number_format = "0.00%"
            elif nome == "Configuracao":
                sheet.column_dimensions["A"].width = 31
                sheet.column_dimensions["B"].width = 125
                for linha in sheet.iter_rows(min_row=2):
                    linha[1].alignment = Alignment(wrap_text=True, vertical="center")
                    sheet.row_dimensions[linha[0].row].height = 32
    return saida


def gerar_graficos(e, pasta):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from statsmodels.graphics.tsaplots import plot_acf, plot_pacf

    destino = Path(pasta) / "resultados_manchas_solares"
    destino.mkdir(exist_ok=True)
    dados = e["dados"]
    fig, axes = plt.subplots(2, 1, figsize=(13, 7))
    axes[0].plot(dados["Ano"], dados["Atividade solar"], color="#243B53", linewidth=1)
    axes[0].axvline(dados['Ano'].iloc[e['fim_treino']], color="#BA6F30", linestyle="--", label="Início da validação")
    axes[0].axvline(dados['Ano'].iloc[e['fim_validacao']], color="#45866C", linestyle="--", label="Início do teste")
    axes[0].set(title="Manchas solares anuais — divisão temporal", ylabel="Atividade solar")
    axes[0].legend()
    teste = e["teste"]
    axes[1].plot(teste["Ano previsto"], teste["Valor real"], label="Real", color="#243B53")
    axes[1].plot(teste["Ano previsto"], teste["Valor predito"], label=f"AR({e['p']}) — um ano à frente", color="#BA6F30")
    axes[1].plot(teste["Ano previsto"], teste["Persistência"], label="Persistência", color="#888888", alpha=0.6)
    axes[1].set(title="Previsões no teste", xlabel="Ano", ylabel="Atividade solar")
    axes[1].legend()
    for ax in axes: ax.grid(True, alpha=0.2)
    fig.tight_layout()
    fig.savefig(destino / "previsoes.png", dpi=160)
    plt.close(fig)
    treino = dados["Atividade solar"].iloc[:e["fim_treino"]]
    fig, axes = plt.subplots(1, 2, figsize=(12, 4))
    plot_acf(treino, lags=30, ax=axes[0], title="Autocorrelação — somente treino")
    plot_pacf(treino, lags=30, method="ywm", ax=axes[1], title="Autocorrelação parcial — somente treino")
    fig.tight_layout()
    fig.savefig(destino / "autocorrelacao.png", dpi=160)
    plt.close(fig)
