"""Busca adicional e auditoria: preservar AR9 quando a melhora global não se confirma."""
import json
import hashlib
from pathlib import Path
import numpy as np
import pandas as pd
from threadpoolctl import threadpool_limits
from openpyxl.styles import Font,PatternFill,Alignment
from manchas_solares import carregar,executar as original,erro_relativo
from manchas_avancados import avaliar_dev,prever_bloco,candidatos
from manchas_decisao import congelar
from investigacao_manchas import metricas,bootstrap_ganho


def assinatura(pasta):
    h=hashlib.sha256()
    for nome in ['manchas_avancados.py','manchas_decisao.py','dados_manchas_solares/sunspots.csv']:
        h.update((Path(pasta)/nome).read_bytes())
    return h.hexdigest()


def executar(pasta,refazer=False,progresso=False):
    pasta=Path(pasta);dados,fonte=carregar(pasta);y=dados['Atividade solar'].to_numpy()
    cache=pasta/'resultados_manchas_solares/cache_avancado';cache.mkdir(parents=True,exist_ok=True)
    ass=assinatura(pasta);manifesto=cache/'manifesto.json'
    valido=(not refazer and manifesto.exists() and json.loads(manifesto.read_text())['sha256']==ass)
    if valido:
        busca=pd.read_csv(cache/'busca.csv');matriz=np.load(cache/'previsoes_dev.npy')
        configs=json.loads((cache/'configs.json').read_text())
        if progresso:print('Busca de desenvolvimento carregada do cache verificado.',flush=True)
    else:
        # O gerador de busca só recebe observações até 1946.
        configs,busca,matriz=avaliar_dev(y[:247],progresso=progresso)
        busca.to_csv(cache/'busca.csv',index=False);np.save(cache/'previsoes_dev.npy',matriz)
        (cache/'configs.json').write_text(json.dumps(configs,indent=2))
        manifesto.write_text(json.dumps(dict(sha256=ass,origem='1700–1946 somente',candidatos=len(configs)),indent=2)+'\n')
    decisao,blocos=congelar(busca,matriz,y[120:247])
    # A decisão da validação é salva antes de calcular resultados de 1947 em diante.
    (cache/'decisao_final_congelada.json').write_text(json.dumps(decisao,indent=2)+'\n')
    blocos.to_csv(cache/'blocos_externos.csv',index=False)
    selecionado=decisao['recomendado']['ids']
    with threadpool_limits(limits=1):
        componentes=np.array([prever_bloco(y,247,len(y),configs[i]) for i in selecionado])
        experimental=componentes.mean(axis=0)
        ar9=prever_bloco(y,247,len(y),configs[0])
    antigo=original(pasta)['teste']['Valor predito'].to_numpy()
    real=y[247:];persistencia=y[246:-1];escala=np.mean(abs(np.diff(y[:247])))
    previsoes=pd.DataFrame({
        'Ano de origem':dados['Ano'].iloc[246:-1].to_numpy(),
        'Ano previsto':dados['Ano'].iloc[247:].to_numpy(),
        'Valor real':real,'Valor predito AR9':ar9,'Taxa de erro AR9 (%)':erro_relativo(real,ar9),
        'Predito experimental':experimental,'Taxa de erro experimental (%)':erro_relativo(real,experimental),
        'Erro absoluto AR9':abs(ar9-real),'Erro absoluto experimental':abs(experimental-real),
        'Persistência':persistencia,'Predito AR19 original':antigo,
    })
    comparacao=pd.DataFrame([
        dict(Modelo=nome,**metricas(real,pred,escala))
        for nome,pred in [('AR9 padrão preservado',ar9),('Candidato experimental',experimental),
                          ('AR19 original',antigo),('Persistência',persistencia)]
    ])
    # Não confundimos promoção na validação com ganho confirmado no período já visto.
    motivos=pd.DataFrame([
        ['Escolha pela validação',decisao['recomendado']['tipo']],
        ['IDs selecionados',json.dumps(selecionado)],
        ['Regra de promoção na validação',decisao['criterio_promocao']['regra']],
        ['Regra satisfeita no desenvolvimento',decisao['criterio_promocao']['promovido']],
        ['Resultado retrospectivo', 'O experimental melhorou MAPE, mas piorou MAE e RMSE frente ao AR9'],
        ['Modelo padrão mantido','AR9 na raiz quadrada, Ridge alpha=1, reajuste anual'],
        ['Protocolo','Todos os modelos usam somente valores reais anteriores, nunca suas próprias previsões'],
        ['Candidatos',len(configs)],
        ['Comparação final','1947–2008 já foi visto. Não é um teste novo. Não houve nova rodada de ajuste após esta comparação'],
        ['Série atualizada','Consulta SILSO v2 realizada; download completo falhou. Nenhum dado parcial ou versão diferente foi incorporado'],
        ['Fonte da base efetivamente usada',fonte['url_csv']],
    ],columns=['Item','Valor'])
    por_faixa=[]
    for a,b,nome in [(-1,10,'Até 10'),(10,50,'10 a 50'),(50,100,'50 a 100'),(100,np.inf,'Acima de 100')]:
        mask=(real>a)&(real<=b)
        for modelo,pred in [('AR9',ar9),('Experimental',experimental)]:
            por_faixa.append(dict(Modelo=modelo,Faixa=nome,Anos=int(mask.sum()),**metricas(real[mask],pred[mask],escala)))
    desempenho_componentes=[]
    for i,pred in zip(selecionado,componentes):
        desempenho_componentes.append(dict(ID=i,Familia=configs[i]['familia'],**metricas(real,pred,escala)))
    configuracoes=pd.DataFrame([dict(ID=i,**c,config=json.dumps(c)) for i,c in enumerate(configs)])
    configuracoes['lags']=configuracoes['lags'].map(json.dumps)
    familias=busca.sort_values(['MAE_cv','p','id']).drop_duplicates('familia')
    return dict(dados=dados,fonte=fonte,configs=configs,busca=busca,matriz=matriz,
                decisao=decisao,blocos=blocos,previsoes=previsoes,comparacao=comparacao,
                motivos=motivos,faixas=pd.DataFrame(por_faixa),familias=familias,
                componentes=pd.DataFrame(desempenho_componentes),configuracoes=configuracoes,
                bootstrap=bootstrap_ganho(real,ar9,experimental))


def exportar(e,pasta):
    destino=Path(pasta)/'resultados_manchas_solares';destino.mkdir(exist_ok=True)
    saida=destino/'investigacao_avancada_manchas.xlsx'
    tabelas={
        'Previsoes':e['previsoes'],'Comparacao':e['comparacao'],
        'Validacao_externa':e['blocos'],'Busca_177_modelos':e['busca'],
        'Melhor_por_familia':e['familias'],'Erros_por_faixa':e['faixas'],
        'Componentes_experimentais':e['componentes'],'Configuracoes':e['configuracoes'],
        'Incerteza_ganho':e['bootstrap'],'Decisao':e['motivos'],
    }
    with pd.ExcelWriter(saida,engine='openpyxl') as writer:
        for nome,tabela in tabelas.items():
            linha_inicio = 2 if nome == 'Previsoes' else 0
            tabela.to_excel(writer,sheet_name=nome,index=False,startrow=linha_inicio)
            sheet=writer.sheets[nome]
            linha_cabecalho=linha_inicio+1
            sheet.freeze_panes=f'A{linha_cabecalho+1}'
            sheet.auto_filter.ref=f'A{linha_cabecalho}:{sheet.cell(sheet.max_row,sheet.max_column).coordinate}'
            sheet.sheet_view.showGridLines=False
            sheet.row_dimensions[linha_cabecalho].height=36
            if nome=='Previsoes':
                inicio=int(tabela['Ano previsto'].min());fim=int(tabela['Ano previsto'].max())
                sheet['A1']=f'Dados de {inicio} até {fim}'
                sheet['A1'].font=Font(name='Arial',size=14,bold=True,color='243B53')
                sheet.row_dimensions[1].height=27
                sheet['A2']='Previsões de teste — comparação retrospectiva; um ano à frente'
                sheet['A2'].font=Font(name='Arial',size=10,italic=True,color='555555')
            for celula in sheet[linha_cabecalho]:
                celula.fill=PatternFill('solid',fgColor='243B53')
                celula.font=Font(name='Arial',size=10,bold=True,color='FFFFFF')
                celula.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
                sheet.column_dimensions[celula.column_letter].width=min(37,max(21,len(str(celula.value))+2))
            for linha in sheet.iter_rows(min_row=linha_cabecalho+1):
                for celula in linha:
                    celula.font=Font(name='Arial',size=10)
                    celula.alignment=Alignment(vertical='center')
                    if isinstance(celula.value,float):celula.number_format='0.000'
                    cab=tabela.columns[celula.column-1]
                    if cab in ['Taxa de erro AR9 (%)','Taxa de erro experimental (%)','MAPE','WAPE','variacao']:celula.number_format='0.00%'
            if nome=='Decisao':
                sheet.column_dimensions['A'].width=42;sheet.column_dimensions['B'].width=115
                for linha in sheet.iter_rows(min_row=linha_cabecalho+1):
                    linha[1].alignment=Alignment(wrap_text=True,vertical='center')
                    sheet.row_dimensions[linha[0].row].height=36
            elif nome=='Configuracoes':
                sheet.column_dimensions['B'].width=25
    return saida


def relatorio(e,pasta):
    m=e['comparacao'].set_index('Modelo');ar=m.loc['AR9 padrão preservado'];novo=m.loc['Candidato experimental']
    decisao=e['decisao'];selecionados=decisao['recomendado']['ids']
    ids=e['familias'][['familia','MAE_cv']]
    boot=e['bootstrap'].iloc[0]
    linhas=[
        '# Investigação avançada de autorregressão',
        '',
        '## Conclusão',
        '',
        'O AR9 continua como padrão. A busca adicional encontrou um candidato melhor na validação e no erro percentual retrospectivo, mas ele piorou MAE e RMSE na comparação histórica. Não há evidência para prometer um modelo globalmente melhor nem um máximo atingido.',
        '',
        '## O que foi testado',
        '',
        f"{len(e['configs'])} configurações adicionais, sem variáveis externas: AR com histórico recente/pesos por recência, seleção esparsa de defasagens, diferenças, interações quadráticas, splines, Kernel Ridge RBF, SVR, modelos por regime de subida/queda ou nível, ElasticNet e gradient boosting.",
        'O horizonte continua sendo um ano à frente. Cada ajuste recebe apenas o histórico real anterior ao alvo. A transformação, a padronização, os regimes e as janelas são calculados apenas nesse histórico. Nenhuma previsão é reinserida como entrada.',
        '',
        '## Seleção antes de consultar 1947–2008 nesta rodada',
        '',
        'Quatro blocos: 1820–1851, 1852–1883, 1884–1915 e 1916–1946. Comparamos o melhor modelo individual com a média dos três melhores de famílias diferentes. O critério é MAE, não MAPE.',
        'Na avaliação externa, cada bloco escolhe seus modelos exclusivamente pelos blocos anteriores. A regra registrada antes da comparação exige pelo menos 2% menos MAE externo que AR9, sem piora maior que 20% em nenhum bloco.',
        f"A regra foi satisfeita: MAE externo {decisao['criterio_promocao']['MAE_externo']:.3f}, contra {decisao['criterio_promocao']['MAE_AR9_externo']:.3f} do AR9. O candidato final é uma média simples dos modelos {selecionados}.",
        '',
        'Componentes escolhidos:',
    ]
    for i in selecionados:linhas.append('- '+json.dumps(e['configs'][i],ensure_ascii=False))
    linhas += ['','Melhor configuração de cada família na própria busca (medida otimista):','','| Família | MAE CV |','| --- | ---: |']
    for _,r in ids.iterrows():linhas.append(f"| {r['familia']} | {r['MAE_cv']:.3f} |")
    linhas += ['','## Comparação retrospectiva de 1947–2008','','| Modelo | MAE | RMSE | R² | MAPE |','| --- | ---: | ---: | ---: | ---: |']
    for nome,r in m.iterrows():linhas.append(f"| {nome} | {r['MAE']:.2f} | {r['RMSE']:.2f} | {r['R2']:.3f} | {r['MAPE']:.2%} |")
    linhas += [
        '',
        f"O candidato apresentou MAE {(novo['MAE']/ar['MAE']-1):.1%} maior e RMSE {(novo['RMSE']/ar['RMSE']-1):.1%} maior. O MAPE caiu de {ar['MAPE']:.2%} para {novo['MAPE']:.2%}. Isso é uma troca entre métricas, não uma melhora geral.",
        f"Ganho de MAE (positivo favorece o candidato), bootstrap descritivo por blocos de 11 anos: {boot['Ganho médio MAE']:.3f}; intervalo percentil de 95% [{boot['Limite 2,5%']:.3f}, {boot['Limite 97,5%']:.3f}]. A análise não corrige pelo amplo processo de seleção e usa poucos ciclos.",
        'A explicação é compatível com limitação de extrapolação de modelos não lineares e mudanças na amplitude dos ciclos; os erros por faixa e os componentes individuais estão no Excel. Essa interpretação é uma hipótese diagnóstica, não uma causa demonstrada.',
        '',
        '## Limites da investigação',
        '',
        'O período de 1947–2008 já era conhecido. A nova busca não recebe esses alvos, mas foi motivada por resultados anteriores: a comparação é retrospectiva. Não houve nova rodada de ajuste depois de abrir seus resultados nesta etapa.',
        'A regra externa também é um critério de seleção do procedimento; seus números não substituem um teste novo. O ganho no desenvolvimento não se transferiu integralmente para o período histórico posterior.',
        'Continuar aumentando a busca no mesmo conjunto pode melhorar números por acaso. Não adotamos um componente isolado só porque ele eventualmente teria um resultado melhor nesse período já visto.',
        'Consultamos a série anual atualizada SILSO v2, mas não foi possível obter o arquivo completo nesta sessão. Nenhum número foi preenchido manualmente, e versões diferentes do índice não foram concatenadas. Sem esse arquivo verificado, a avaliação em anos novos ficou pendente.',
        '',
        '## Executar',
        '',
        '```bash',
        'uv run --locked python prever_manchas.py --avancado',
        '# Refazer todos os 177 candidatos:',
        'uv run --locked python prever_manchas.py --avancado --refazer-busca',
        '```',
        '',
        'O comando padrão sem --avancado mantém o AR9 e a investigação anterior. A busca usa um cache com hash do código e do CSV para garantir correspondência. O Excel avançado exibe os dois resultados lado a lado, sem trocar silenciosamente o modelo padrão.',
        '',
        '## Fontes técnicas',
        '',
        '- https://scikit-learn.org/stable/modules/generated/sklearn.kernel_ridge.KernelRidge.html',
        '- https://scikit-learn.org/stable/modules/generated/sklearn.preprocessing.SplineTransformer.html',
        '- https://www.sidc.be/SILSO/infosnytot',
    ]
    p=Path(pasta)/'resultados_manchas_solares/relatorio_avancado.md';p.write_text('\n'.join(linhas)+'\n');return p


def graficos(e,pasta):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    r=e['previsoes']
    fig,axes=plt.subplots(2,1,figsize=(13,8))
    for coluna,label,cor in [('Valor real','Real','#243B53'),('Valor predito AR9','AR9 mantido','#45866C'),('Predito experimental','Experimental','#BA6F30')]:
        axes[0].plot(r['Ano previsto'],r[coluna],label=label,color=cor)
    axes[0].set(title='Comparação retrospectiva — candidato não supera AR9 em todas as métricas',ylabel='Atividade solar')
    axes[0].legend()
    b=e['blocos'];x=np.arange(len(b))
    axes[1].bar(x-.18,b['MAE_AR9'],width=.36,label='AR9',color='#45866C')
    axes[1].bar(x+.18,b['MAE'],width=.36,label='Procedimento candidato',color='#BA6F30')
    axes[1].set(xticks=x,xticklabels=[f'{a}–{b}' for a,b in zip(b['inicio'],b['fim'])],ylabel='MAE',title='Avaliação externa antes de 1947')
    axes[1].legend()
    for ax in axes:ax.grid(True,alpha=.2)
    fig.tight_layout();fig.savefig(Path(pasta)/'resultados_manchas_solares/comparacao_avancada.png',dpi=160);plt.close(fig)
